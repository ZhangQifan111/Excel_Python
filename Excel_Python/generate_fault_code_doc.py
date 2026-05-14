#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
故障码对应关系文档生成器
生成Word文档，用于后期维护时查看规则
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime

def set_cell_shading(cell, color):
    """设置单元格背景色"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def create_fault_code_document():
    """创建故障码对应关系文档"""
    
    doc = Document()
    
    # 设置标题
    title = doc.add_heading('户储光伏平台产品故障码对应关系', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 文档信息
    info_para = doc.add_paragraph()
    info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_para.add_run(f'版本: V1.00\n').font.size = Pt(10)
    info_para.add_run(f'生成日期: {datetime.now().strftime("%Y-%m-%d")}\n').font.size = Pt(10)
    info_para.add_run('用途: 故障数据分析时故障码解析参考').font.size = Pt(10)
    
    doc.add_paragraph()
    
    # ===== 第一部分：故障码编号规则 =====
    doc.add_heading('一、故障码编号规则', level=1)
    
    rule_para = doc.add_paragraph()
    rule_para.add_run('故障码格式: ').bold = True
    rule_para.add_run('E + 分类编号(1位) + 序号(2位)\n\n')
    rule_para.add_run('示例:\n')
    rule_para.add_run('• E100 → 分类1(电网) + 序号00\n')
    rule_para.add_run('• E405 → 分类4(电池) + 序号05\n')
    
    doc.add_paragraph()
    
    # ===== 第二部分：寄存器地址对应关系 =====
    doc.add_heading('二、寄存器地址对应关系', level=1)
    
    # 创建表格
    table = doc.add_table(rows=7, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 表头
    headers = ['故障码范围', '寄存器地址', '寄存器名称', '故障类型']
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True
        set_cell_shading(cell, '4472C4')  # 蓝色背景
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)  # 白色字体
    
    # 数据行
    data = [
        ('E100 ~ E108', '20712', '电网故障（用户展示）', '电网相关'),
        ('E200 ~ E202', '20713', '离网侧故障（用户展示）', '离网侧相关'),
        ('E300 ~ E306', '20714', 'PV侧故障（用户展示）', 'PV侧相关'),
        ('E400 ~ E412', '20715', '电池故障（用户展示）', '电池相关'),
        ('E500 ~ E501', '20716', '油机故障（用户展示）', '油机相关'),
        ('E600 ~ E609', '20717', '系统故障（用户展示）', '系统相关'),
    ]
    
    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, cell_data in enumerate(row_data):
            table.rows[row_idx].cells[col_idx].text = cell_data
    
    doc.add_paragraph()
    
    # ===== 第三部分：故障码与bit位对应规则 =====
    doc.add_heading('三、故障码与bit位对应规则', level=1)
    
    rule_para2 = doc.add_paragraph()
    rule_para2.add_run('核心规则:\n').bold = True
    rule_para2.add_run('• E100 = 寄存器20712的bit0\n')
    rule_para2.add_run('• E101 = 寄存器20712的bit1\n')
    rule_para2.add_run('• E200 = 寄存器20713的bit0\n')
    rule_para2.add_run('• ...\n\n')
    rule_para2.add_run('规律: ').bold = True
    rule_para2.add_run('故障码后两位数字 = bit位编号\n')
    rule_para2.add_run('例如: E405 → 寄存器20715的bit5')
    
    doc.add_paragraph()
    
    # ===== 第四部分：详细故障码列表 =====
    doc.add_heading('四、详细故障码列表', level=1)
    
    # 定义所有故障码
    fault_codes = {
        '电网故障 (寄存器20712)': [
            ('E100', 'bit0', '电网未接', 'No AC Connection', '告警'),
            ('E101', 'bit1', '电网电压异常', 'Grid voltage abnormal', '告警'),
            ('E102', 'bit2', '电网频率异常', 'Grid frequency abnormal', '告警'),
            ('E103', 'bit3', '电网相序接反', 'Grid phase sequence is error', '告警'),
            ('E104', 'bit4', '输出DCI超范围', 'Output DCI over-range', '告警'),
            ('E105', 'bit5', '防逆流超限故障', 'Anti-backflow over-limit fault', '告警'),
            ('E106', 'bit6', '电表或CT异常', 'Meter or CT abnormal', '告警'),
            ('E107', 'bit7', '电表或CT反接', 'Meter or CT reverse connection', '告警'),
            ('E108', 'bit8', '电网未满足并网范围', 'Grid Connection requirements not met', '告警'),
        ],
        '离网侧故障 (寄存器20713)': [
            ('E200', 'bit0', '离网输出电压过高', 'The off-grid output voltage is too high', '告警'),
            ('E201', 'bit1', '离网输出电压过低', 'The off-grid output voltage is too low', '告警'),
            ('E202', 'bit2', '交流过载', 'Ac overload', '告警'),
        ],
        'PV侧故障 (寄存器20714)': [
            ('E300', 'bit0', 'PV输入过压', 'PV input overvoltage', '告警'),
            ('E301', 'bit1', 'PV组串反接', 'PV string reverse connection', '告警'),
            ('E302', 'bit2', 'PV过载', 'PV overload', '告警'),
            ('E303', 'bit3', '拉弧故障', 'AFCI fault', '告警'),
            ('E304', 'bit4', 'GFCI异常', 'Residual current High', '告警'),
            ('E305', 'bit5', '绝缘阻抗低', 'PV Isolation Low', '告警'),
            ('E306', 'bit6', 'PV短路', 'PV Short Circuit', '告警'),
        ],
        '电池故障 (寄存器20715)': [
            ('E400', 'bit0', '电池BMS故障', 'Battery BMS fault', '故障'),
            ('E401', 'bit1', '电池BMS运行异常', 'Battery BMS running abnormal', '告警'),
            ('E402', 'bit2', '电池BMS通信异常', 'Battery BMS communication abnormal', '故障'),
            ('E403', 'bit3', '电池DCDC故障', 'Battery DCDC fault', '故障'),
            ('E404', 'bit4', '电池DCDC运行异常', 'Battery DCDC run abnormal', '告警'),
            ('E405', 'bit5', '电池超额输出', 'Battery Overdischarge', '告警'),
            ('E406', 'bit6', '电池过压', 'Battery overvoltage', '告警'),
            ('E407', 'bit7', '电池欠压', 'Battery undervoltage', '告警'),
            ('E408', 'bit8', '电池电量低', 'Battery SOC is low', '提示'),
            ('E409', 'bit9', '电池过温故障', 'Battery Over-temperature protection', '告警'),
            ('E410', 'bit10', '电池低温故障', 'Battery Low-temperature protection', '告警'),
            ('E411', 'bit11', '电池端口过压故障', 'Battery Port overvoltage', '告警'),
            ('E412', 'bit12', '电池加热异常', 'Battery PTC Abnormal', '告警'),
        ],
        '油机故障 (寄存器20716)': [
            ('E500', 'bit0', '油机输出过载', 'Oil machine overload', '告警'),
            ('E501', 'bit1', '油机故障', 'Oil machine fault', '故障'),
        ],
        '系统故障 (寄存器20717)': [
            ('E600', 'bit0', '内部通信故障', 'Internal communication fault', '故障'),
            ('E601', 'bit1', 'PCS运行异常', 'PCS is running abnormally', '告警'),
            ('E602', 'bit2', 'PCS故障', 'PCS fault', '故障'),
            ('E603', 'bit3', 'OTA升级失败', 'OTA upgrade failure', '故障'),
            ('E604', 'bit4', '过温保护', 'Over-temperature protection', '告警'),
            ('E605', 'bit5', '风扇异常', 'Fan abnormal', '故障'),
            ('E606', 'bit6', '并机异常', 'Parallel machine abnormal', '故障'),
            ('E607', 'bit7', '系统版本不一致', 'System Version Incompatible', '告警'),
            ('E608', 'bit8', '软件版本错误', 'Software version wrong', '故障'),
            ('E609', 'bit9', '软件参数有误', 'Software parameter incorrect', '告警'),
        ],
    }
    
    # 为每个故障类型创建表格
    for category, codes in fault_codes.items():
        doc.add_heading(category, level=2)
        
        # 创建表格
        table = doc.add_table(rows=len(codes)+1, cols=5)
        table.style = 'Table Grid'
        
        # 表头
        headers = ['故障码', 'bit位', '中文描述', '英文描述', '类型']
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True
            set_cell_shading(cell, '4472C4')
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        # 数据行
        for row_idx, code_data in enumerate(codes, start=1):
            for col_idx, cell_data in enumerate(code_data):
                cell = table.rows[row_idx].cells[col_idx]
                cell.text = str(cell_data)
                # 根据类型设置颜色
                if col_idx == 4:  # 类型列
                    if cell_data == '故障':
                        set_cell_shading(cell, 'FFCCCC')  # 浅红色
                    elif cell_data == '告警':
                        set_cell_shading(cell, 'FFFFCC')  # 浅黄色
                    elif cell_data == '提示':
                        set_cell_shading(cell, 'CCFFCC')  # 浅绿色
        
        doc.add_paragraph()
    
    # ===== 第五部分：数据解析示例 =====
    doc.add_heading('五、数据解析示例', level=1)
    
    example_para = doc.add_paragraph()
    example_para.add_run('示例数据:\n').bold = True
    example_para.add_run('电池故障（用户展示）(20715) = [2, 4]\n\n')
    
    example_para.add_run('解析步骤:\n').bold = True
    example_para.add_run('1. 识别寄存器地址: 20715 → 电池故障\n')
    example_para.add_run('2. 解析bit位: [2, 4] → bit2 和 bit4\n')
    example_para.add_run('3. 查表转换:\n')
    example_para.add_run('   • bit2 → E402 → 电池BMS通信异常\n')
    example_para.add_run('   • bit4 → E404 → 电池DCDC运行异常\n\n')
    
    example_para.add_run('解析结果:\n').bold = True
    example_para.add_run('该时刻存在故障: 电池BMS通信异常、电池DCDC运行异常')
    
    doc.add_paragraph()
    
    # ===== 第六部分：故障分类规则 =====
    doc.add_heading('六、故障分类规则', level=1)
    
    classify_para = doc.add_paragraph()
    classify_para.add_run('1. 故障: ').bold = True
    classify_para.add_run('产品出现功能失效的异常，上报故障\n')
    classify_para.add_run('2. 告警: ').bold = True
    classify_para.add_run('产品出现功能未失效的异常，仍可自动恢复的，上报告警。产品外部环境条件不满足的，也上报告警。\n')
    classify_para.add_run('3. 提示: ').bold = True
    classify_para.add_run('产品运行正常会出现的状况，想提示给客户看的，像电池电量低的，上报提示')
    
    doc.add_paragraph()
    
    # ===== 第七部分：内部故障码映射关系 =====
    doc.add_heading('七、内部故障码映射关系（参考）', level=1)
    
    note_para = doc.add_paragraph()
    note_para.add_run('说明: ').bold = True
    note_para.add_run('内部故障码（寄存器20650-20696）与用户展示故障码（寄存器20712-20717）之间存在映射关系。\n')
    note_para.add_run('详细映射关系请参考《平台逆变器错误代码对照表.xlsx》中的"映射关系"列。\n\n')
    note_para.add_run('示例映射:\n')
    note_para.add_run('• 20650:bit0 (电网断电) → E100 (电网未接)\n')
    note_para.add_run('• 20650:bit1 (电网安规过压保护) → E101 (电网电压异常)\n')
    note_para.add_run('• 20659:bit0 (电池1预充故障) → E400 (电池BMS故障)\n')
    
    # ===== 页脚 =====
    doc.add_paragraph()
    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.add_run('— 文档结束 —').font.size = Pt(9)
    
    # 保存文档
    output_path = '故障码对应关系文档.docx'
    doc.save(output_path)
    print(f'文档已生成: {output_path}')
    
    return output_path

if __name__ == '__main__':
    create_fault_code_document()
